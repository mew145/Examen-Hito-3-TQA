# -*- coding: utf-8 -*-
"""
Examen Practico Hito 3 - TQA
Rellena la plantilla unificada de casos de prueba (Plantilla_TC_Unificada_TQA.docx)
con los 5 casos de caja negra: Parte 1 (diseno) y Parte 2 (ejecucion real).

Entradas:
  - docs/Plantilla_TC_Unificada_TQA.docx  (copia de la plantilla de la materia)
  - scripts/resultados.json               (resultados reales de ejecucion)
  - gen_informe.DISENO                    (diseno de los 5 casos)

Salida:
  - docs/casos_caja_negra_plantilla_llena.docx

Uso: python scripts/gen_plantilla_llena.py
"""

import copy
import json
import pathlib
import sys

import docx

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from gen_informe import DISENO  # noqa: E402

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "docs" / "Plantilla_TC_Unificada_TQA.docx"
DATOS = RAIZ / "scripts" / "resultados.json"
SALIDA = RAIZ / "docs" / "casos_caja_negra_plantilla_llena.docx"

EJECUTADO_POR = "Yuri Jesús (GitHub: mew145)"


def set_cell(cell, texto):
    """Reemplaza el contenido de una celda conservando su formato de parrafo."""
    if not cell.paragraphs:
        cell.text = texto
        return
    parrafo = cell.paragraphs[0]
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for extra in parrafo.runs[1:]:
            extra.text = ""
    else:
        parrafo.add_run(texto)
    # limpiar parrafos adicionales de la celda
    for p in cell.paragraphs[1:]:
        p.text = ""


def fill_diseno(tbl, tc, res):
    filas = tbl.rows
    set_cell(filas[1].cells[1], tc["id"])
    set_cell(filas[1].cells[3], tc["prioridad"])
    set_cell(filas[2].cells[1], tc["titulo"])
    set_cell(filas[2].cells[3], tc["tipo"])
    set_cell(filas[3].cells[1], tc["modulo"])
    set_cell(filas[3].cells[3], "QA (https://www.saucedemo.com/)")
    set_cell(filas[4].cells[0], "¿QUÉ SE QUIERE VALIDAR?  " + tc["validar"])
    set_cell(filas[5].cells[1], tc["estado_inicial"])
    set_cell(filas[6].cells[1], tc["usuario"])
    set_cell(filas[7].cells[1], tc["config"])

    # filas de datos de prueba (fila 8 = primer dato con 4 columnas)
    datos = tc["datos"]
    set_cell(filas[8].cells[1], datos[0][0])
    set_cell(filas[8].cells[2], datos[0][1])
    set_cell(filas[8].cells[3], datos[0][2])
    resto = datos[1:]
    set_cell(filas[9].cells[0], "Otros datos")
    if resto:
        texto = "; ".join(f"{n}: {v} ({o})" for n, v, o in resto)
        set_cell(filas[9].cells[1], texto)
    else:
        set_cell(filas[9].cells[1], "Sin datos adicionales")
    set_cell(filas[10].cells[0], "Origen")
    set_cell(filas[10].cells[1], "Generado por el tester / ambiente de demostración")
    set_cell(filas[11].cells[0], "Observación")
    set_cell(filas[11].cells[1], "Usuario y contrasena oficiales publicados en la propia pagina de login")

    # pasos: filas 13..19
    for i in range(7):
        fila = filas[13 + i]
        if i < len(tc["pasos"]):
            paso, dato, esperado = tc["pasos"][i]
            set_cell(fila.cells[1], paso)
            set_cell(fila.cells[2], dato)
            set_cell(fila.cells[3], esperado)
        else:
            set_cell(fila.cells[1], "")
            set_cell(fila.cells[2], "")
            set_cell(fila.cells[3], "")
    set_cell(filas[20].cells[1], tc["criterio"])


def fill_ejecucion(tbl, tc, res):
    filas = tbl.rows
    estado = res.get("estado", "NOT RUN")
    set_cell(filas[1].cells[1], tc["id"])
    set_cell(filas[2].cells[1], EJECUTADO_POR)
    set_cell(filas[3].cells[1], res.get("fecha", "-"))
    set_cell(filas[4].cells[1], "QA (https://www.saucedemo.com/)")
    set_cell(filas[4].cells[3], "Build en linea del sitio de demostracion")
    casillas = "    ".join(
        f"{'[X]' if estado == e else '[ ]'} {e}" for e in ("PASS", "FAIL", "BLOCKED", "NOT RUN")
    )
    set_cell(filas[5].cells[1], casillas)
    set_cell(filas[6].cells[1], res.get("observacion", "-"))
    coincidencia = ("Coincide con lo esperado; el comportamiento observado es el correcto."
                    if estado == "PASS" else "No coincide con lo esperado; ver observacion y defecto.")
    set_cell(filas[7].cells[1], coincidencia)

    pasos = res.get("pasos", [])
    for i in range(7):
        fila = filas[9 + i]
        if i < len(pasos):
            p = pasos[i]
            set_cell(fila.cells[1], p.get("paso", ""))
            set_cell(fila.cells[2], "PASS" if estado == "PASS" else "FAIL")
            set_cell(fila.cells[3], str(p.get("obtenido", "")))
        else:
            set_cell(fila.cells[1], "")
            set_cell(fila.cells[2], "")
            set_cell(fila.cells[3], "")

    capturas = ", ".join(res.get("capturas", [])) or "N/A"
    set_cell(filas[16].cells[1],
             f"Tipo: Captura    |    Referencia: {capturas} (carpeta evidencias/capturas/)    |    "
             "Descripción: demuestra el comportamiento observado durante la ejecución")
    if estado == "PASS":
        set_cell(filas[17].cells[1], "Desviación encontrada: N/A (la prueba coincidió con lo esperado).    "
                                     "Referencia del defecto: N/A")
    else:
        set_cell(filas[17].cells[1], "Desviación encontrada: ver documentación.    "
                                     "Referencia del defecto: BUG-001 (ver bugs/bugs_encontrados.md)")
    set_cell(filas[18].cells[1],
             "Ejecución automatizada con Playwright (scripts/ejecutar_pruebas.py); "
             "capturas en evidencias/capturas/.")


def main():
    if not PLANTILLA.exists():
        raise SystemExit(f"No se encuentra la plantilla: {PLANTILLA}")
    resultados = {c["id"]: c for c in json.loads(DATOS.read_text(encoding="utf-8"))["casos"]}

    doc = docx.Document(str(PLANTILLA))
    tbl_diseno_orig = doc.tables[0]
    tbl_ejec_orig = doc.tables[1]
    cuerpo = doc.element.body

    # TC-001 rellena las tablas originales de la plantilla
    tc1 = DISENO[0]
    fill_diseno(tbl_diseno_orig, tc1, resultados.get(tc1["id"], {}))
    fill_ejecucion(tbl_ejec_orig, tc1, resultados.get(tc1["id"], {}))

    # TC-002 .. TC-005: copias profundas de ambas tablas
    for tc in DISENO[1:]:
        res = resultados.get(tc["id"], {})
        doc.add_page_break()
        doc.add_heading(f"{tc['id']} · {tc['titulo']}", level=2)
        doc.add_paragraph("PARTE 1 · DISEÑO DEL CASO DE PRUEBA").runs[0].bold = True
        copia_diseno = copy.deepcopy(tbl_diseno_orig._tbl)
        cuerpo.append(copia_diseno)
        doc.add_paragraph("PARTE 2 · EJECUCIÓN DEL CASO DE PRUEBA").runs[0].bold = True
        copia_ejec = copy.deepcopy(tbl_ejec_orig._tbl)
        cuerpo.append(copia_ejec)
        # rellenar las copias recien agregadas
        fill_diseno(doc.tables[-2], tc, res)
        fill_ejecucion(doc.tables[-1], tc, res)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(SALIDA))
    print(f"Plantilla generada: {SALIDA}")


if __name__ == "__main__":
    main()
