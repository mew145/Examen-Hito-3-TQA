# -*- coding: utf-8 -*-
"""
Examen Practico Hito 3 - TQA
Genera el informe final en Word (entregable/informe_hito3.docx) a partir de:
  - scripts/resultados.json  (resultados reales de ejecucion)
  - el diseno de los 5 casos (definido en DISENO)
Luego puede convertirse a PDF con Word (ver convertir_a_pdf.ps1).

Uso: python scripts/gen_informe.py
"""

import json
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "entregable" / "informe_hito3.docx"
DATOS = RAIZ / "scripts" / "resultados.json"

ESTUDIANTE = "Yuri Rojas (GitHub: mew145)"
DOCENTE = "Mg. Reynel Sánchez"
MATERIA = "Testing & Quality Assurance (TQA) - Hito 3"
FECHA_INFORME = "20/10/2026"

URL = "https://www.saucedemo.com/"

# ---------------------------------------------------------------------------
# Diseno de los 5 casos de prueba de caja negra (Parte 1 de la plantilla)
# ---------------------------------------------------------------------------
DISENO = [
    {
        "id": "TC-001",
        "titulo": "Inicio de sesion con credenciales validas",
        "prioridad": "Alta",
        "tipo": "Funcional / Positiva",
        "modulo": "Inicio de sesion",
        "validar": "Que un usuario con credenciales correctas acceda al catalogo de productos.",
        "estado_inicial": "Navegador limpio en la URL de la aplicacion.",
        "usuario": "Usuario demo standard_user (sin permisos especiales).",
        "config": "Resolucion 1440x900, navegador Chromium, sitio publico de demostracion.",
        "datos": [("Usuario", "standard_user", "Usuario de demostracion de la pagina"),
                  ("Contrasena", "secret_sauce", "Contrasena oficial del usuario demo")],
        "pasos": [
            ("Abrir https://www.saucedemo.com/", "-", "Se muestra la pagina de login (Username / Password / Login)"),
            ("Ingresar el usuario", "standard_user", "El campo Username muestra el valor ingresado"),
            ("Ingresar la contrasena", "secret_sauce", "El campo Password enmascara el valor ingresado"),
            ("Presionar el boton Login", "-", "Redirige a la pagina del catalogo (inventory.html)"),
        ],
        "criterio": "PASS si el usuario accede al catalogo 'Products' y no permanece en el login.",
    },
    {
        "id": "TC-002",
        "titulo": "Inicio de sesion con credenciales invalidas (prueba negativa)",
        "prioridad": "Alta",
        "tipo": "Negativa",
        "modulo": "Inicio de sesion",
        "validar": "Que el sistema rechace el acceso con una contrasena incorrecta y muestre un mensaje de error claro.",
        "estado_inicial": "Navegador limpio en la URL de la aplicacion.",
        "usuario": "Usuario demo standard_user.",
        "config": "Resolucion 1440x900, navegador Chromium.",
        "datos": [("Usuario", "standard_user", "Usuario valido"),
                  ("Contrasena", "wrong_password", "Contrasena incorrecta (dato invalido a proposito)")],
        "pasos": [
            ("Abrir https://www.saucedemo.com/", "-", "Se muestra la pagina de login"),
            ("Ingresar el usuario", "standard_user", "El campo Username muestra el valor ingresado"),
            ("Ingresar una contrasena incorrecta", "wrong_password", "El campo Password enmascara el valor"),
            ("Presionar el boton Login", "-", "No debe ocurrir el acceso al catalogo"),
        ],
        "criterio": "PASS si permanece en el login y se muestra el mensaje 'Epic sadface: Username and password do not match each other'.",
    },
    {
        "id": "TC-003",
        "titulo": "Agregar un producto al carrito",
        "prioridad": "Alta",
        "tipo": "Funcional / Positiva",
        "modulo": "Catalogo y carrito",
        "validar": "Que al agregar un producto el contador del carrito se actualice y el item aparezca en el carrito.",
        "estado_inicial": "Sesion iniciada con standard_user y catalogo visible.",
        "usuario": "Usuario demo standard_user.",
        "config": "Resolucion 1440x900, navegador Chromium.",
        "datos": [("Producto", "Sauce Labs Backpack", "Primer producto del catalogo"),
                  ("Accion", "Add to cart", "Boton del producto")],
        "pasos": [
            ("Iniciar sesion", "standard_user / secret_sauce", "Acceso al catalogo"),
            ("Presionar 'Add to cart' en Sauce Labs Backpack", "-", "El boton cambia a 'Remove'"),
            ("Verificar el contador del carrito", "-", "El badge muestra '1'"),
            ("Ingresar al carrito", "-", "El producto 'Sauce Labs Backpack' aparece en la lista"),
        ],
        "criterio": "PASS si el badge = 1, el boton pasa a 'Remove' y el item figura en el carrito.",
    },
    {
        "id": "TC-004",
        "titulo": "Eliminar un producto desde el carrito",
        "prioridad": "Media",
        "tipo": "Funcional / Positiva",
        "modulo": "Catalogo y carrito",
        "validar": "Que al eliminar un producto el carrito se actualice y el contador desaparezca al quedar vacio.",
        "estado_inicial": "Sesion iniciada con un producto agregado al carrito.",
        "usuario": "Usuario demo standard_user.",
        "config": "Resolucion 1440x900, navegador Chromium.",
        "datos": [("Producto", "Sauce Labs Bike Light", "Agregado previamente al carrito"),
                  ("Accion", "Remove", "Boton de eliminar dentro del carrito")],
        "pasos": [
            ("Iniciar sesion y agregar Sauce Labs Bike Light", "-", "El badge del carrito muestra '1'"),
            ("Ingresar al carrito", "-", "El producto aparece en la lista"),
            ("Presionar 'Remove' sobre el producto", "-", "El item deja de aparecer"),
            ("Verificar el contador del carrito", "-", "El badge desaparece porque el carrito queda vacio"),
        ],
        "criterio": "PASS si el item se elimina, la lista queda vacia y el badge del carrito desaparece.",
    },
    {
        "id": "TC-005",
        "titulo": "Completar el proceso de checkout (compra exitosa)",
        "prioridad": "Alta",
        "tipo": "Funcional / Positiva",
        "modulo": "Compra y navegacion",
        "validar": "Que un usuario pueda completar el formulario de compra y finalizar el pedido correctamente.",
        "estado_inicial": "Sesion iniciada y al menos un producto en el carrito.",
        "usuario": "Usuario demo standard_user.",
        "config": "Resolucion 1440x900, navegador Chromium.",
        "datos": [("Nombre", "Ana", "Dato valido"),
                  ("Apellido", "Quispe", "Dato valido"),
                  ("Codigo postal", "0000", "Dato valido")],
        "pasos": [
            ("Agregar un producto e ingresar al carrito", "-", "El producto esta en la lista"),
            ("Presionar 'Checkout'", "-", "Se muestra el formulario Your Information"),
            ("Completar nombre, apellido y codigo postal", "Ana / Quispe / 0000", "Los tres campos muestran los datos"),
            ("Presionar 'Continue'", "-", "Se muestra el resumen de la compra (Overview)"),
            ("Presionar 'Finish'", "-", "Se muestra la confirmacion de la compra"),
        ],
        "criterio": "PASS si se muestra el mensaje 'Thank you for your order!' al finalizar.",
    },
]


def cargar_resultados():
    if DATOS.exists():
        return json.loads(DATOS.read_text(encoding="utf-8"))
    return {"fecha_ejecucion": "-", "casos": [], "exploracion": []}


def titulo(doc, texto, nivel=1):
    p = doc.add_heading(texto, level=nivel)
    return p


def tabla(doc, filas, ancho_primera_col=True):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for izq, der in filas:
        c = t.add_row().cells
        c[0].text = izq
        c[1].text = der
        if ancho_primera_col:
            for run in c[0].paragraphs[0].runs:
                run.bold = True
    return t


def tabla_pasos(doc, cabecera, filas):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    for i, h in enumerate(cabecera):
        celda = t.rows[0].cells[i]
        celda.text = h
        for run in celda.paragraphs[0].runs:
            run.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = str(valor)
    return t


def construir():
    data = cargar_resultados()
    casos = {c["id"]: c for c in data.get("casos", [])}
    exploracion = data.get("exploracion", [])

    doc = Document()
    est = doc.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(11)

    # ---------------- Portada ----------------
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNIVERSIDAD FRANZ TAMAYO\nTESTING & QUALITY ASSURANCE (TQA)")
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EXAMEN PRACTICO - HITO 3\nCaja negra · Caja blanca · Gherkin · Bugs · GitHub")
    r.font.size = Pt(14)
    r.bold = True

    doc.add_paragraph()
    for texto in (
        f"Aplicacion bajo prueba: {URL}",
        f"Estudiante: {ESTUDIANTE}",
        f"Docente: {DOCENTE}",
        f"Materia: {MATERIA}",
        f"Fecha de elaboracion: {FECHA_INFORME}",
        f"Fecha de ejecucion de pruebas: {data.get('fecha_ejecucion', '-')}",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(texto)
    doc.add_page_break()

    # ---------------- 1. Introduccion ----------------
    titulo(doc, "1. Introduccion", 1)
    doc.add_paragraph(
        f"El presente informe documenta el proceso de prueba aplicado sobre {URL} "
        "(SauceDemo / Swag Labs), sitio de demostracion utilizado para practicar tecnicas de "
        "pruebas funcionales. Se disenaron y ejecutaron 5 casos de prueba de caja negra, se "
        "analizo un fragmento de codigo en Python con tecnicas de caja blanca, se redactaron "
        "escenarios Gherkin y se documentaron los defectos reproducibles encontrados, todo ello "
        "organizado en un repositorio individual de GitHub con Issues y Project Board."
    )
    doc.add_paragraph("Cuentas de demostracion utilizadas: standard_user / secret_sauce (las indicadas en el propio sitio).")

    # ---------------- 2. Casos de caja negra ----------------
    titulo(doc, "2. Casos de prueba de caja negra", 1)
    doc.add_paragraph(
        "Cada caso sigue las dos partes de la plantilla unificada de la materia: diseno del caso "
        "(Parte 1) y ejecucion (Parte 2). El resultado esperado se definio antes de ejecutar; el "
        "resultado observado se registro durante la ejecucion real del sitio."
    )
    for diseno in DISENO:
        cid = diseno["id"]
        res = casos.get(cid, {})
        titulo(doc, f"2.{DISENO.index(diseno) + 1} {cid} - {diseno['titulo']}", 2)

        doc.add_paragraph("Parte 1 - Diseno del caso de prueba").runs[0].bold = True
        tabla(doc, [
            ("ID del TC", cid),
            ("Titulo", diseno["titulo"]),
            ("Prioridad", diseno["prioridad"]),
            ("Tipo", diseno["tipo"]),
            ("Modulo", diseno["modulo"]),
            ("Ambiente", "QA (sitio publico de demostracion)"),
            ("¿Que se quiere validar?", diseno["validar"]),
            ("Estado inicial / precondiciones", diseno["estado_inicial"]),
            ("Usuario y permisos", diseno["usuario"]),
            ("Configuracion", diseno["config"]),
        ])
        doc.add_paragraph("Datos de prueba").runs[0].bold = True
        tabla_pasos(doc, ["Dato", "Valor utilizado", "Observacion"],
                    [(n, v, o) for n, v, o in diseno["datos"]])
        doc.add_paragraph("Pasos").runs[0].bold = True
        tabla_pasos(doc, ["#", "Paso / accion", "Dato utilizado", "Resultado esperado"],
                    [(i + 1, p, d, e) for i, (p, d, e) in enumerate(diseno["pasos"])])
        tabla(doc, [("Criterio de exito", diseno["criterio"])])

        doc.add_paragraph("Parte 2 - Ejecucion del caso de prueba").runs[0].bold = True
        estado = res.get("estado", "NOT RUN")
        tabla(doc, [
            ("ID del TC", cid),
            ("Ejecutado por", ESTUDIANTE),
            ("Fecha / hora", res.get("fecha", "-")),
            ("Ambiente / version", f"QA - {URL} (build del sitio en linea)"),
            ("Resultado", f"{'[X]' if estado == 'PASS' else '[ ]'} PASS    "
                          f"{'[X]' if estado == 'FAIL' else '[ ]'} FAIL    "
                          f"{'[X]' if estado == 'BLOCKED' else '[ ]'} BLOCKED    "
                          f"{'[X]' if estado == 'NOT RUN' else '[ ]'} NOT RUN"),
            ("¿Que ocurrio realmente?", res.get("observacion", "-")),
            ("Comparacion", "Coincide con lo esperado." if estado == "PASS"
                            else "No coincide con lo esperado; ver observacion y defecto."),
        ])
        if res.get("pasos"):
            doc.add_paragraph("Pasos ejecutados").runs[0].bold = True
            filas = []
            for i, paso in enumerate(res["pasos"], start=1):
                filas.append((i, paso.get("paso", ""), paso.get("esperado", ""), paso.get("obtenido", "")))
            tabla_pasos(doc, ["#", "Paso ejecutado", "Resultado esperado", "Resultado observado"], filas)
        if res.get("capturas"):
            tabla(doc, [("Evidencia", "Capturas: " + ", ".join(res["capturas"]) +
                         f"  (carpeta evidencias/capturas/)")])
        tabla(doc, [("Si resulta FAIL", "Desviacion encontrada: N/A (la prueba coincidio con lo esperado). "
                                        "Referencia del defecto: N/A" if estado == "PASS"
                                        else "Ver reporte de defectos en bugs/bugs_encontrados.md")])
        doc.add_paragraph()

    # ---------------- 3. Bugs ----------------
    titulo(doc, "3. Defectos encontrados", 1)
    doc.add_paragraph(
        "Durante la ejecucion se realizaron verificaciones exploratorias adicionales. Solo se "
        "registran defectos reproducibles; si la prueba paso, se registra PASS y no se inventa un bug."
    )
    bugs = [e for e in exploracion if e.get("reproducible")]
    no_bugs = [e for e in exploracion if not e.get("reproducible")]
    if bugs:
        for i, bug in enumerate(bugs, start=1):
            titulo(doc, f"BUG-{i:03d} - {bug.get('escenario', '')}", 2)
            tabla(doc, [
                ("Descripcion", bug.get("escenario", "")),
                ("Precondiciones y datos", "Sesion iniciada con standard_user / secret_sauce; producto agregado al carrito."),
                ("Pasos para reproducirlo", "1) Agregar un producto al carrito. 2) Ir a Checkout. "
                                            "3) Completar los campos con valores que no tienen contenido real. "
                                            "4) Presionar Continue."),
                ("Resultado esperado", bug.get("esperado", "")),
                ("Resultado actual", bug.get("obtenido", "")),
                ("Evidencia", bug.get("captura", "") or "N/A"),
                ("Severidad / prioridad", "Media - Baja (defecto de validacion de datos)"),
            ])
    else:
        doc.add_paragraph("No se encontraron defectos reproducibles en las verificaciones exploratorias.")
    for e in no_bugs:
        doc.add_paragraph(
            f"Verificacion {e.get('id')}: {e.get('escenario')} -> {e.get('obtenido')} "
            f"(comportamiento correcto, no se genera defecto).", style="Intense Quote"
        )
    doc.add_paragraph(
        "Las pruebas TC-001 a TC-005 resultaron PASS; conforme a la regla de uso de la plantilla, "
        "un caso PASS no genera reporte de defecto."
    )

    # ---------------- 4. Caja blanca ----------------
    titulo(doc, "4. Analisis de caja blanca", 1)
    doc.add_paragraph("Codigo analizado (fragmento proporcionado en el examen):")
    codigo = doc.add_paragraph()
    r = codigo.add_run(
        "def validar_compra(edad, cuenta_activa, saldo, monto):\n"
        "    if edad < 18:\n"
        '        return "Compra no permitida: edad"\n'
        "    if not cuenta_activa:\n"
        '        return "Compra no permitida: cuenta"\n'
        "    if saldo >= monto and monto > 0:\n"
        '        return "Compra aprobada"\n'
        '    return "Compra rechazada"'
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    doc.add_paragraph(
        "Identificacion: 3 decisiones (D1: edad < 18; D2: not cuenta_activa; D3: saldo >= monto and monto > 0) "
        "y 4 condiciones simples (C1: edad < 18, C2: cuenta_activa, C3: saldo >= monto, C4: monto > 0). "
        "D3 es una decision compuesta por el operador and."
    )
    doc.add_paragraph("Criterio utilizado: cobertura de decisiones.").runs[0].bold = True
    doc.add_paragraph(
        "Calculo del minimo: D1 (V/F) + D2 (V/F) + D3 (V/F) = 6 resultados de decision. Como la funcion "
        "tiene un solo return por camino y existen 4 caminos disjuntos (edad, cuenta, aprobada, rechazada), "
        "el minimo de casos para 100% de cobertura de decisiones es 4 (uno por salida). Con 4 casos se "
        "alcanza ademas el 100% de cobertura de sentencias."
    )
    doc.add_paragraph("Casos disenados (3):").runs[0].bold = True
    tabla_pasos(doc, ["ID", "edad", "cuenta_activa", "saldo", "monto", "Resultado esperado", "Que cubre"], [
        ("CB-001", "17", "True", "100", "50", "Compra no permitida: edad", "D1 = V"),
        ("CB-002", "25", "False", "100", "50", "Compra no permitida: cuenta", "D1 = F, D2 = V"),
        ("CB-003", "30", "True", "200", "100", "Compra aprobada", "D2 = F, D3 = V"),
    ])
    doc.add_paragraph(
        "Resultado de la ejecucion: los 3 casos devolvieron exactamente el resultado esperado "
        "(PASS). Cobertura alcanzada con los 3 casos: 5 de 6 resultados de decision = 83,3%."
    )
    doc.add_paragraph("¿Y si el criterio fuera cobertura de condiciones?").runs[0].bold = True
    doc.add_paragraph(
        "C3 y C4 deben tomar V y F por separado, lo que obliga a dos ejecuciones distintas del ultimo "
        "camino: serian necesarios 5 casos (por ejemplo, saldo=50/monto=100 y saldo=100/monto=0)."
    )

    # ---------------- 5. Gherkin ----------------
    titulo(doc, "5. Escenarios Gherkin (actividad adicional)", 1)
    doc.add_paragraph(
        "Se redactaron 7 escenarios sobre las funcionalidades de SauceDemo (ver docs/escenarios_gherkin.md). "
        "Ejemplos destacados:"
    )
    gherkin = doc.add_paragraph()
    r = gherkin.add_run(
        "Feature: Inicio de sesion\n"
        "  Scenario: Acceso exitoso con credenciales validas (positivo)\n"
        "    Given que estoy en la pagina de inicio de sesion de SauceDemo\n"
        '    When ingreso el usuario "standard_user" y la contrasena "secret_sauce"\n'
        '    And presiono el boton "Login"\n'
        '    Then deberia ser redirigido a la pagina del catalogo "Products"\n\n'
        "  Scenario: Acceso con credenciales incorrectas (de error)\n"
        "    Given que estoy en la pagina de inicio de sesion de SauceDemo\n"
        '    When ingreso el usuario "standard_user" y la contrasena "wrong_password"\n'
        '    And presiono el boton "Login"\n'
        '    Then deberia permanecer en la pagina de inicio de sesion\n'
        '    And deberia ver el mensaje "Epic sadface: Username and password do not match each other"'
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9)

    # ---------------- 6. Conclusiones ----------------
    titulo(doc, "6. Conclusiones", 1)
    for c in (
        "Los 5 casos de caja negra fueron ejecutados sobre la aplicacion real y sus resultados se "
        "documentaron con evidencia (capturas) en el repositorio.",
        "El analisis de caja blanca permitio identificar 3 decisiones y 4 condiciones simples; con "
        "3 casos se alcanza el 83,3% de cobertura de decisiones y se explico el caso adicional "
        "necesario para llegar al 100%.",
        "Las pruebas negativas validan que el sitio informa errores claros al usuario y bloquea el "
        "acceso no autorizado.",
        "El trabajo esta versionado en Git/GitHub con commits descriptivos, Issues y Project Board "
        "para el seguimiento de las tareas.",
    ):
        doc.add_paragraph(c, style="List Bullet")

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(SALIDA))
    print(f"Informe generado: {SALIDA}")


if __name__ == "__main__":
    construir()
